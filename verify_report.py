"""Check exported PDF text, page layout, required sections and Word diagram."""
from pathlib import Path
import json
import argparse
import re
import pymupdf
from docx import Document
from PIL import Image, ImageOps, ImageDraw

root=Path(__file__).resolve().parent
parser=argparse.ArgumentParser()
parser.add_argument('--vietnamese',action='store_true')
args=parser.parse_args()
out=root/'deliverables'
if args.vietnamese:
    out=out/'tieng_viet'
profile=json.loads((root/'student.json').read_text(encoding='utf-8'))
pdfs=list(out.glob('DDM501_Assignment1_*.pdf'))
assert pdfs, 'No submission PDF generated'
pdf=max(pdfs,key=lambda f:f.stat().st_mtime)
report=pymupdf.open(pdf)
text='\n'.join(p.get_text() for p in report)
required=['1. Problem definition (20%)','2. Requirements analysis (20%)','3. Goals and metrics (20%)','4. High-level architecture design (25%)','5. Trade-offs analysis (15%)','References and requirement traceability']
if args.vietnamese:
    required=['1. Định nghĩa bài toán (20%)','2. Phân tích yêu cầu (20%)','3. Mục tiêu và chỉ số đánh giá (20%)','4. Thiết kế kiến trúc tổng thể (25%)','5. Phân tích các đánh đổi (15%)','Tài liệu tham khảo và đối chiếu yêu cầu']
    assert '\ufffd' not in text and '\u25a0' not in text, 'Invalid Unicode glyphs in Vietnamese PDF'
for phrase in required:
    assert phrase in text, f'Missing required section: {phrase}'
issues=[]
page_summary=[]
thumbs=[]
for i,page in enumerate(report):
    blocks=page.get_text('blocks')
    body=[b for b in blocks if b[1]<page.rect.height-58 and b[4].strip()]
    if not body:
        issues.append({'page':i+1,'reason':'Blank body page'})
    for b in body:
        if b[0]<50 or b[2]>page.rect.width-50 or b[3]>page.rect.height-52:
            issues.append({'page':i+1,'bounds':list(b[:4]),'text':b[4][:80]})
    page_summary.append({'page':i+1,'first_line':page.get_text().splitlines()[0],'words':len(page.get_text().split()),'body_bottom':round(max((b[3] for b in body),default=0),1)})
    pix=page.get_pixmap(matrix=pymupdf.Matrix(.55,.55))
    im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    im=ImageOps.expand(im,border=5,fill='#cbd5e1')
    thumbs.append(im)
cols=3
rows=(len(thumbs)+cols-1)//cols
w=max(im.width for im in thumbs)
hh=max(im.height for im in thumbs)+22
sheet=Image.new('RGB',(cols*w,rows*hh),'#e5eaf0')
draw=ImageDraw.Draw(sheet)
for i,im in enumerate(thumbs):
    x=(i%cols)*w;y=(i//cols)*hh
    sheet.paste(im,(x,y));draw.text((x+8,y+im.height+3),f'Page {i+1}',fill='black')
sheet.save(out/'layout_preview.png')
doc=Document(pdf.with_suffix('.docx'))
assert len(doc.inline_shapes)==1, 'Editable Word report must contain architecture diagram'
if args.vietnamese:
    cover_text=report[0].get_text()
    word_text='\n'.join(p.text for p in doc.paragraphs)
    for value in profile.values():
        assert value in cover_text, 'Student information missing from PDF cover'
        assert value in word_text, 'Student information missing from Word report'
    source=(out/'REPORT.md').read_text(encoding='utf-8')
    assert '{{' not in source, 'Unresolved placeholders in final report'
    body_text='\n'.join(line for page in report for line in page.get_text().splitlines()
                        if not line.startswith('DDM501 |') and not re.fullmatch(r'\d+',line.strip()))
    compact_pdf=re.sub(r'\s+','',body_text)
    # Verify prose/heading coverage independently of page and table layouts.
    for line in source.splitlines():
        line=line.strip()
        if not line or line.startswith(('|','![','---')):
            continue
        content=re.sub(r'^#+\s*','',line)
        assert re.sub(r'\s+','',content) in compact_pdf, 'Source text missing from PDF: '+ascii(content[:80])
    if 'Phiên bản 2.' in source:
        defined=set(re.findall(r'^\[(\d+)\]',source,flags=re.M))
        cited=set(re.findall(r'\[(\d+)\]',source))
        assert cited<=defined, 'Citation without reference entry'
        assert {str(i) for i in range(1,12)}<=defined, 'Missing upgraded academic references'
        assert report.get_toc(), 'PDF bookmarks missing'
    if 'Phiên bản 2.1' in source:
        assert 'Mục lục' in text and 'Danh mục từ viết tắt' in text, 'Missing academic front matter'
        assert 'Hình 1.' in text, 'Architecture caption missing'
        assert 'Bảng 1.' in text, 'Numbered table captions missing'
        for level,label,page in report.get_toc():
            if re.match(r'^[1-5]\. ',label) or re.match(r'^[1-5]\.\d',label):
                assert re.sub(r'\s+','',label) in re.sub(r'\s+','',report[page-1].get_text()), 'Bookmark page mismatch'
result={'pdf':pdf.name,'pages':len(report),'words':len(text.split()),'required_sections':'PASS','word_architecture':'PASS','layout_issues':issues,'page_summary':page_summary,'identity_complete':profile['student_id']!='StudentID' and profile['name']!='Name','status':'PASS' if not issues else 'REVIEW'}
(out/'verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
assert not issues, 'Layout boundary issues found'
