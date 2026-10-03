"""Export the Vietnamese Markdown source to PDF and editable Word."""
from pathlib import Path
import json
import re
from xml.sax.saxutils import escape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak, Flowable, Spacer
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'deliverables'/'tieng_viet'
OUT.mkdir(parents=True,exist_ok=True)
profile=json.loads((ROOT/'student.json').read_text(encoding='utf-8'))
safe=lambda value:re.sub(r'[<>:"/\\|?*\s]+','_',value).strip('_')
stem=f"DDM501_Assignment1_{safe(profile['student_id'])}_{safe(profile['name'])}"
fontdir=Path('C:/Windows/Fonts')
for name,file in [('Arial','arial.ttf'),('Arial-Bold','arialbd.ttf'),('Arial-Italic','ariali.ttf'),('Arial-BoldItalic','arialbi.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(fontdir/file)))
pdfmetrics.registerFontFamily('Arial',normal='Arial',bold='Arial-Bold',italic='Arial-Italic',boldItalic='Arial-BoldItalic')
for name,file in [('TimesVN','times.ttf'),('TimesVN-Bold','timesbd.ttf'),('TimesVN-Italic','timesi.ttf'),('TimesVN-BoldItalic','timesbi.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(fontdir/file)))
pdfmetrics.registerFontFamily('TimesVN',normal='TimesVN',bold='TimesVN-Bold',italic='TimesVN-Italic',boldItalic='TimesVN-BoldItalic')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='VNBody',fontName='TimesVN',fontSize=12,leading=17.5,spaceAfter=8,alignment=TA_JUSTIFY,textColor=colors.HexColor('#243247')))
styles.add(ParagraphStyle(name='VNCell',fontName='TimesVN',fontSize=10,leading=13,spaceAfter=0))
styles.add(ParagraphStyle(name='CoverTitle',fontName='TimesVN-Bold',fontSize=23,leading=31,spaceAfter=22,alignment=TA_CENTER,textColor=colors.HexColor('#123e58')))
styles.add(ParagraphStyle(name='CoverMeta',fontName='TimesVN',fontSize=12,leading=19,spaceAfter=12,alignment=TA_CENTER))
for name,size,leading in [('Heading1',17,23),('Heading2',11.5,16),('Heading3',10.5,15)]:
    styles[name].fontName='TimesVN-Bold'
    styles[name].fontSize=size
    styles[name].leading=leading
    styles[name].textColor=colors.HexColor('#123e58' if name=='Heading1' else '#146b73')
doc=Document()
section=doc.sections[0]
section.top_margin=Inches(.7)
section.bottom_margin=Inches(.7)
section.header.paragraphs[0].text='DDM501 | Bài tập cá nhân 1 | Thiết kế hệ thống ML'
doc.styles['Normal'].font.name='Times New Roman'
doc.styles['Normal'].font.size=Pt(12)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].font.name='Times New Roman'
    doc.styles[name].paragraph_format.keep_with_next=True
section.page_width=Inches(A4[0]/72)
section.page_height=Inches(A4[1]/72)
section.left_margin=Inches(56/72)
section.right_margin=Inches(56/72)
section.different_first_page_header_footer=True
page_field=OxmlElement('w:fldSimple');page_field.set(qn('w:instr'),'PAGE')
section.footer.paragraphs[0]._p.append(page_field)
section.footer.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
story=[]
bookmark_count=0
table_count=0

class ReportDocument(SimpleDocTemplate):
    def afterFlowable(self,flowable):
        if hasattr(flowable,'report_heading'):
            label,level,key=flowable.report_heading
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(label,key,level=level,closed=False)
            if re.match(r'^[1-5]\. ',label) or re.match(r'^[1-5]\.\d',label) or label.startswith(('Tài liệu tham khảo','Đối chiếu yêu cầu')):
                toc_level=1 if re.match(r'^[1-5]\.\d',label) else 0
                self.notify('TOCEntry',(toc_level,label,self.page,key))

class Architecture(Flowable):
    def __init__(self):
        super().__init__();self.width=483;self.height=320
    def draw(self):
        c=self.canv
        def box(x,y,w,hh,lines,fill='#e6f0f4'):
            c.setFillColor(colors.HexColor(fill));c.setStrokeColor(colors.HexColor('#6b8a9a'))
            c.roundRect(x,y,w,hh,5,stroke=1,fill=1)
            c.setFillColor(colors.HexColor('#19384d'));c.setFont('Arial',8)
            for i,line in enumerate(lines):c.drawCentredString(x+w/2,y+hh-14-11*i,line)
        def arrow(x1,y1,x2,y2,label=None):
            import math
            c.setStrokeColor(colors.HexColor('#496a7c'));c.setLineWidth(1);c.line(x1,y1,x2,y2)
            a=math.atan2(y2-y1,x2-x1)
            for offset in [-.45,.45]:c.line(x2,y2,x2-6*math.cos(a+offset),y2-6*math.sin(a+offset))
            if label:c.setFont('Arial',7);c.drawString((x1+x2)/2+3,(y1+y2)/2+3,label)
        box(5,262,137,48,['POS / tồn kho / ERP','Sản phẩm / khuyến mãi'])
        box(173,262,137,48,['Thu thập và kiểm tra ngày','Lưu nguồn bất biến'])
        box(341,262,137,48,['Đặc trưng đã xử lý','Ảnh chụp có phiên bản'])
        arrow(142,286,173,286);arrow(310,286,341,286)
        box(341,172,137,53,['Huấn luyện / đánh giá','Một LightGBM dùng chung','Kiểm tra ứng viên'])
        arrow(409,262,409,225,'hằng tuần')
        box(173,172,137,53,['Kho mô hình MLflow','Có người duyệt phát hành','Cố định phiên bản'])
        arrow(341,198,310,198)
        box(5,172,137,53,['Suy luận theo lô ngày','Dự báo / quy tắc đặt hàng','Kiểm tra / dự phòng'])
        arrow(173,198,142,198)
        arrow(341,270,329,270);arrow(329,270,329,242);arrow(329,242,73,242);arrow(73,242,73,225)
        box(5,82,137,53,['Công bố trọn bộ nguyên tử','Lưu kết quả PostgreSQL','API chỉ đọc'])
        arrow(73,172,73,135)
        box(173,82,137,53,['Giao diện lập kế hoạch','Phê duyệt / điều chỉnh','Xuất đơn nháp sang ERP'])
        arrow(142,108,173,108)
        box(341,82,137,53,['Giám sát và nhật ký','Thực tế / phân bố / SLO','Cảnh báo / yêu cầu xét lại'])
        arrow(310,108,341,108);arrow(409,135,409,172,'qua kiểm tra')
        box(173,5,305,47,['Phản hồi: lượng bán thực tế, hết hàng, kết quả đơn, điều chỉnh','Kiểm tra vận hành ngày; rà soát chất lượng tuần'],fill='#f0f4e8')
        arrow(241,82,241,52);arrow(409,52,409,82)
        c.setFont('Arial-Italic',7);c.setFillColor(colors.HexColor('#536b80'))
        c.drawString(5,20,'Mũi tên: dữ liệu / điều khiển')

diagram_pdf=OUT/'architecture.pdf'
c=canvas.Canvas(str(diagram_pdf),pagesize=(503,340));c.translate(10,10);Architecture().drawOn(c,0,0);c.save()
with pymupdf.open(diagram_pdf) as render:
    render[0].get_pixmap(matrix=pymupdf.Matrix(2,2)).save(str(OUT/'architecture.png'))

def text(s):
    story.append(Paragraph(escape(s),styles['VNBody']))
    paragraph=doc.add_paragraph(s);paragraph.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
def heading(s,level):
    global bookmark_count
    bookmark_count+=1
    title=Paragraph(escape(s),styles[f'Heading{level}'])
    title.report_heading=(s,level-1,f'heading-{bookmark_count}')
    story.append(title);doc.add_heading(s,level=level)
def table(lines):
    global table_count
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows=[rows[0]]+rows[2:]
    n=len(rows[0])
    if rows[0][0]!='Viết tắt':
        table_count+=1
        labels={'Phần bắt buộc':'Cấu trúc báo cáo theo rubric','Bên liên quan':'Các bên liên quan và yêu cầu','Nguồn nghiên cứu':'Tổng hợp cơ sở học thuật','Mã':'Yêu cầu chức năng và nghiệm thu','Khía cạnh':'Yêu cầu phi chức năng','Tầng và chỉ số':'Mục tiêu và chỉ số đánh giá','Ngưỡng đề xuất':'Cơ sở ngưỡng và cách kiểm chứng','Thời gian':'Lịch vận hành hằng ngày','Thành phần, công cụ xem xét':'Trách nhiệm các thành phần hệ thống','Tình huống':'Phương án xử lý sự cố','Phương án':'So sánh ứng viên dự báo','Bằng chứng cần lưu':'Bằng chứng phát hành','Đánh đổi':'Các đánh đổi thiết kế','Yêu cầu trong PDF':'Đối chiếu yêu cầu đề bài'}
        caption=f'Bảng {table_count}. '+labels.get(rows[0][0],'Tổng hợp thiết kế')+'.'
        story.append(Paragraph(escape(caption),styles['VNCell']))
        doc.add_paragraph(caption)
    widths=[483/n]*n
    if n==3:
        if rows[0][0]=='Mã':widths=[43,230,210]
        elif rows[0][1]=='Tỷ trọng':widths=[245,60,178]
        elif rows[0][0]=='Đánh đổi':widths=[97,193,193]
        else:widths=[125,166,192]
    elif n==2:
        widths=[100,383] if rows[0][0]=='Thời gian' else [180,303]
    cells=[[Paragraph(escape(val),styles['VNCell']) for val in row] for row in rows]
    t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dcecf1')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#8fb1bd')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f7fa')]),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#dfe5eb'))]))
    t.spaceAfter=8;story.append(t)
    dt=doc.add_table(rows=1,cols=n);dt.style='Light Shading Accent 1'
    for cell,val in zip(dt.rows[0].cells,rows[0]):cell.text=val
    for row in rows[1:]:
        for cell,val in zip(dt.add_row().cells,row):cell.text=val
    for row in dt.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:run.font.size=Pt(10)
    for cell in dt.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:run.bold=True
    for row in dt.rows[1:]:
        cant_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant_split)
    repeat=OxmlElement('w:tblHeader');dt.rows[0]._tr.get_or_add_trPr().append(repeat)

source=(ROOT/'REPORT_VI.md').read_text(encoding='utf-8')
for key,val in profile.items():source=source.replace('{{'+key+'}}',val)
lines=source.splitlines()
summary_index=lines.index('## Tóm tắt')
story.append(Spacer(1,45))
for line in lines[:summary_index]:
    line=line.strip()
    if not line:continue
    if line.startswith('#'):
        title=re.sub(r'^#+\s*','',line)
        story.append(Paragraph(escape(title),styles['CoverTitle']))
        paragraph=doc.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=paragraph.add_run(title);run.bold=True;run.font.size=Pt(22)
    else:
        story.append(Paragraph(escape(line),styles['CoverMeta']))
        paragraph=doc.add_paragraph(line);paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
story.append(PageBreak());doc.add_page_break()
heading('Mục lục',1)
toc=TableOfContents()
toc.levelStyles=[ParagraphStyle(name='TOCMain',fontName='TimesVN-Bold',fontSize=12,leading=17,spaceBefore=7),ParagraphStyle(name='TOCSub',fontName='TimesVN',fontSize=11,leading=15,leftIndent=16,firstLineIndent=0,spaceBefore=2)]
story.append(toc)
toc_paragraph=doc.add_paragraph()
toc_field=OxmlElement('w:fldSimple');toc_field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
for line in lines:
    label=re.sub(r'^#+\s*','',line.strip())
    if line.startswith('#') and (re.match(r'^[1-5]\.\d',label) or re.match(r'^[1-5]\. ',label) or label.startswith(('Tài liệu tham khảo','Đối chiếu yêu cầu'))):
        run=OxmlElement('w:r');val=OxmlElement('w:t');val.text=label;run.append(val)
        br=OxmlElement('w:br');run.append(br);toc_field.append(run)
toc_paragraph._p.append(toc_field)
story.append(PageBreak());doc.add_page_break()
heading('Danh mục từ viết tắt',1)
glossary=[['ML','Machine Learning – học máy'],['MVP','Minimum Viable Product – phiên bản khả dụng tối thiểu'],['POS / ERP','Điểm bán / hệ thống hoạch định nguồn lực doanh nghiệp'],['SKU','Stock Keeping Unit – mã đơn vị hàng hóa'],['API / CI','Giao diện lập trình ứng dụng / tích hợp liên tục'],['ETS','Error, Trend, Seasonality – mô hình sai số, xu hướng, mùa vụ'],['WAPE','Weighted Absolute Percentage Error – sai số phần trăm tuyệt đối có trọng số'],['MAE / MASE','Sai số tuyệt đối trung bình / sai số tuyệt đối trung bình chuẩn hóa'],['DOS','Days of Supply – số ngày tồn kho đủ cung ứng'],['IP / SS','Inventory Position / Safety Stock – vị thế tồn kho / tồn kho an toàn'],['ETA','Estimated Time of Arrival – thời điểm hàng đến dự kiến'],['SLO','Service Level Objective – mục tiêu mức dịch vụ'],['RTO / RPO','Thời gian khôi phục mục tiêu / mức mất dữ liệu tối đa tính theo thời gian'],['PSI','Population Stability Index – chỉ số ổn định quần thể'],['ITT','Intention-to-Treat – phân tích theo nhóm được phân ban đầu']]
table(['| Viết tắt | Ý nghĩa |','| --- | --- |']+['| '+' | '.join(row)+' |' for row in glossary])
story.append(PageBreak());doc.add_page_break()
heading('Tóm tắt và phạm vi báo cáo',1)
i=summary_index
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='---':
        # Start major rubric sections on a new page, allow subsections to flow.
        following=next((s.strip() for s in lines[i+1:] if s.strip()),'')
        if re.match(r'^# [1-5]\. ',following) or following.startswith('# Tài liệu tham khảo') or following.startswith('# Đối chiếu'):
            story.append(PageBreak());doc.add_page_break()
    elif line=='{{architecture}}':
        story.append(Architecture());doc.add_picture(str(OUT/'architecture.png'),width=Inches(6.2))
        caption='Hình 1. Kiến trúc hệ thống dự báo và bổ sung hàng đề xuất. Nguồn: thiết kế của báo cáo.'
        story.append(Paragraph(escape(caption),styles['VNCell']));doc.add_paragraph(caption)
    elif line.startswith('# '):
        subsection=bool(re.match(r'^# [1-5]\.\d',line))
        heading(line[2:],2 if subsection else 1)
    elif line.startswith('## '):heading(line[3:],2)
    elif line.startswith('|'):
        tab=[]
        while i<len(lines) and lines[i].strip().startswith('|'):tab.append(lines[i]);i+=1
        table(tab);continue
    else:
        paragraph=[line]
        while i+1<len(lines) and lines[i+1].strip() and not lines[i+1].startswith(('#','|')) and lines[i+1].strip()!='---':
            i+=1;paragraph.append(lines[i].strip())
        text(' '.join(paragraph))
    i+=1

def footer(c,d):
    if d.page==1:return
    c.saveState();c.setStrokeColor(colors.HexColor('#cbd9e1'));c.line(56,43,A4[0]-56,43)
    c.setFont('Arial',8);c.setFillColor(colors.HexColor('#536b80'))
    c.drawString(56,30,'DDM501 | Dự báo nhu cầu | Thiết kế hệ thống ML')
    c.drawRightString(A4[0]-56,30,str(d.page));c.restoreState()

pdf=OUT/(stem+'.pdf')
ReportDocument(str(pdf),pagesize=A4,rightMargin=56,leftMargin=56,topMargin=48,bottomMargin=58,title='DDM501 – Thiết kế hệ thống dự báo nhu cầu',author=profile['name']).multiBuild(story,onFirstPage=footer,onLaterPages=footer)
doc.save(OUT/(stem+'.docx'))
(OUT/'REPORT.md').write_text(source.replace('{{architecture}}','![Sơ đồ kiến trúc](architecture.png)'),encoding='utf-8')
(OUT/'architecture.mmd').write_text('''flowchart LR
  Nguon[POS / Tồn kho / ERP / Kế hoạch khuyến mãi] --> ThuThap[Thu thập và kiểm tra hằng ngày]
  ThuThap --> Tho[Lưu nguồn bất biến]
  Tho --> DacTrung[Đặc trưng có phiên bản]
  DacTrung --> HuanLuyen[Huấn luyện và đánh giá theo thời gian]
  HuanLuyen --> Duyet[Kiểm tra và duyệt phát hành]
  Duyet --> Kho[MLflow: phiên bản cố định]
  DacTrung --> SuyLuan[Suy luận theo lô hằng ngày]
  Kho --> SuyLuan
  SuyLuan --> QuyTac[Kiểm tra dự báo và quy tắc bổ sung]
  QuyTac --> CSDL[Công bố nguyên tử / PostgreSQL]
  CSDL --> GiaoDien[API đọc và giao diện nhân viên]
  GiaoDien --> PheDuyet[Phê duyệt / điều chỉnh]
  PheDuyet --> ERP[Đơn nháp ERP chống trùng]
  Nguon --> GiamSat[Giám sát chất lượng / thực tế / phân bố]
  SuyLuan --> GiamSat
  PheDuyet --> GiamSat
  GiamSat --> XemXet[Điều tra / yêu cầu huấn luyện lại]
  XemXet --> HuanLuyen
''',encoding='utf-8')
print('Created Vietnamese PDF and Word: '+ascii(pdf.name))
