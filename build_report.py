"""Build the assignment report from a single editable content source."""
from pathlib import Path
import json
import re
from xml.sax.saxutils import escape
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Flowable
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches, Pt
import pymupdf

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
PROFILE = ROOT / 'student.json'
profile = json.loads(PROFILE.read_text(encoding='utf-8')) if PROFILE.exists() else {'student_id':'StudentID', 'name':'Name'}
def safe(s):
    return re.sub(r'[<>:"/\\|?*\s]+', '_', s).strip('_')
stem = f"DDM501_Assignment1_{safe(profile['student_id'])}_{safe(profile['name'])}"
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyReport', fontName='Helvetica', fontSize=9.5, leading=14, spaceAfter=8, textColor=colors.HexColor('#243247')))
styles.add(ParagraphStyle(name='CellReport', fontName='Helvetica', fontSize=8, leading=10.5, spaceAfter=2))
styles.add(ParagraphStyle(name='TitleReport', fontName='Helvetica-Bold', fontSize=27, leading=34, textColor=colors.HexColor('#123e58'), spaceAfter=22))
styles.add(ParagraphStyle(name='SubtitleReport', fontSize=12, leading=18, textColor=colors.HexColor('#536b80'), spaceAfter=14))
styles['Heading1'].fontSize=18
styles['Heading1'].leading=23
styles['Heading1'].textColor=colors.HexColor('#123e58')
styles['Heading2'].fontSize=11.5
styles['Heading2'].leading=16
styles['Heading2'].textColor=colors.HexColor('#146b73')
story=[]
md=[]
doc=Document()
sec=doc.sections[0]
sec.header.paragraphs[0].text='DDM501 | Individual Assignment 1 | ML System Design'
sec.top_margin=Inches(.75)
sec.bottom_margin=Inches(.7)
doc.styles['Normal'].font.name='Calibri'
doc.styles['Normal'].font.size=Pt(10)

def p(text):
    story.append(Paragraph(escape(text), styles['BodyReport']))
    md.append(text+'\n')
    doc.add_paragraph(text)
def h(text,level=2):
    story.append(Paragraph(escape(text),styles['Heading'+str(level)]))
    md.append('#'*level+' '+text+'\n')
    doc.add_heading(text,level=level)
def page(title):
    if story:
        story.append(PageBreak());doc.add_page_break();md.append('\n---\n')
    h(title,1)
def table(headers,rows,widths=None):
    data=[[Paragraph(escape(str(x)),styles['CellReport']) for x in headers]]
    data.extend([[Paragraph(escape(str(x)),styles['CellReport']) for x in row] for row in rows])
    t=Table(data,colWidths=widths or [483/len(headers)]*len(headers),repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dcecf1')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#8fb1bd')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f7fa')]),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#dfe5eb'))]))
    t.spaceAfter=8
    story.append(t)
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |')
    md.extend('| '+' | '.join(str(x).replace('|','/') for x in row)+' |' for row in rows)
    md.append('')
    dt=doc.add_table(rows=1,cols=len(headers));dt.style='Light Shading Accent 1'
    for cell,val in zip(dt.rows[0].cells,headers):cell.text=str(val)
    for row in rows:
        for cell,val in zip(dt.add_row().cells,row):cell.text=str(val)

class Architecture(Flowable):
    def __init__(self):
        super().__init__()
        self.width=483
        self.height=320
    def draw(self):
        c=self.canv
        def box(x,y,w,hh,lines,fill='#e6f0f4'):
            c.setFillColor(colors.HexColor(fill));c.setStrokeColor(colors.HexColor('#6b8a9a'));c.roundRect(x,y,w,hh,5,stroke=1,fill=1)
            c.setFillColor(colors.HexColor('#19384d'));c.setFont('Helvetica',8)
            for i,line in enumerate(lines):c.drawCentredString(x+w/2,y+hh-14-11*i,line)
        def arrow(x1,y1,x2,y2,label=None):
            c.setStrokeColor(colors.HexColor('#496a7c'));c.setLineWidth(1);c.line(x1,y1,x2,y2)
            import math
            a=math.atan2(y2-y1,x2-x1)
            for offset in [-.45,.45]:c.line(x2,y2,x2-6*math.cos(a+offset),y2-6*math.sin(a+offset))
            if label:c.setFont('Helvetica',7);c.drawString((x1+x2)/2+3,(y1+y2)/2+3,label)
        box(5,262,137,48,['POS / inventory / ERP','Product / promotion plans'])
        box(173,262,137,48,['Daily ingestion + validation','Immutable raw storage'])
        box(341,262,137,48,['Curated daily features','Versioned Parquet snapshots'])
        arrow(142,286,173,286);arrow(310,286,341,286)
        box(341,172,137,53,['Training + time backtests','Single global LightGBM','Candidate evaluation'])
        arrow(409,262,409,225,'weekly')
        box(173,172,137,53,['MLflow model registry','Human promotion gate','Pinned approved version'])
        arrow(341,198,310,198)
        box(5,172,137,53,['Daily batch inference','Forecast + rule-based order','Data quality / fallback'])
        arrow(173,198,142,198)
        arrow(341,270,329,270);arrow(329,270,329,242);arrow(329,242,73,242);arrow(73,242,73,225)
        box(5,82,137,53,['Atomic forecast publication','PostgreSQL result store','Read-only API'])
        arrow(73,172,73,135)
        box(173,82,137,53,['Planner dashboard','Approve / override','ERP draft order export'])
        arrow(142,108,173,108)
        box(341,82,137,53,['Monitoring + audit','Actuals / drift / SLOs','Alerts + retrain request'])
        arrow(310,108,341,108)
        arrow(409,135,409,172,'gated trigger')
        box(173,5,305,47,['Feedback: mature sales, stockouts, order outcomes, overrides','Daily operational checks; weekly quality review'],fill='#f0f4e8')
        arrow(241,82,241,52);arrow(409,52,409,82)
        c.setFont('Helvetica-Oblique',7);c.setFillColor(colors.HexColor('#536b80'));c.drawString(5,20,'Solid arrows: data / control flow')

page('ML System Design Document')
story.append(Spacer(1,28))
story.append(Paragraph('Daily Demand Forecasting<br/>for a Retail Chain',styles['TitleReport']))
md.append('# Daily Demand Forecasting for a Retail Chain\n')
doc.add_heading('Daily Demand Forecasting for a Retail Chain',0)
p('Course: DDM501 - AI in DevOps, DataOps, MLOps')
p('Assignment: Individual Assignment 1 | Submission: Written report | Assignment weight: 5%')
p(f"Student ID: {profile['student_id']} | Student name: {profile['name']}")
p('Prepared: 3 October 2026')
h('Executive summary')
p('This document designs a production ML system for a hypothetical Vietnamese grocery retailer with 20 stores and 2,000 active stocked products per store. One shared demand forecasting model produces daily predictions for the next seven days. A deterministic replenishment policy turns forecasts into draft order quantities; planners retain approval authority. The design covers data quality, reliable daily delivery, time-aware evaluation, monitoring, rollback and measurable business outcomes.')
h('Assumptions and evidence status')
p('The retailer, historical baseline values, costs, volumes and acceptance thresholds are design assumptions, not observations from a real organization. No model has been trained and no improvement is claimed as achieved. Baselines must be measured in discovery; proposed targets must be confirmed in a controlled pilot. Official technical documentation is cited for tool capabilities only.')
h('Report structure and assessment alignment')
table(['Required section','Weight','Location'],[['Problem definition','20%','Section 1'],['Requirements analysis','20%','Section 2'],['Goals and metrics','20%','Section 3'],['High-level architecture design','25%','Section 4'],['Trade-offs analysis','15%','Section 5']],[245,60,178])

page('1. Problem definition (20%)')
h('1.1 Context and background')
p('Assume a regional retailer operating 20 grocery stores. Each store stocks approximately 2,000 active packaged food, beverage and household products. Suppliers deliver daily or every few days; replenishment lead time is one to three days. Local buying patterns vary by store, weekday, promotions and holidays. Store managers currently use spreadsheets and ERP stock reports to estimate orders. Under-ordering causes empty shelves and lost purchases; over-ordering ties up cash and increases waste on short shelf-life products.')
p('The scope is daily store-product forecasting for established products. The initial pilot excludes products with less than 56 days of history, highly intermittent items and products requiring specialized expiry optimization. Such items use explicit rule-based fallbacks. The system does not optimize prices, recommend products or introduce interacting predictive models.')
h('1.2 Measurable problem statement')
p('The design objective is to reduce avoidable stockouts and excess inventory by supplying seven-day forecasts before the daily ordering cutoff. Assume an eligible-product baseline stockout rate of 8% of open store-product-days, inventory of 12 days of supply, and four planner hours per store each week. The pilot targets stockouts at or below 6.4% (a 20% relative reduction), at least 10% less inventory at comparable service levels, and planner effort at or below 2.5 hours per store per week. These outcomes must be measured, not inferred solely from forecast accuracy.')
h('1.3 Current situation')
p('Each planner uses a 28-day sales average, manually adjusts for promotions and multiplies by supplier lead time. A separate spreadsheet holds safety stock and pack sizes. There is no shared evaluation process, model version history or systematic feedback from overrides. Store-specific knowledge helps detect unusual events, but inconsistent spreadsheets make planning difficult to audit. This current policy is the business control arm. A same-weekday four-week average is also tested as a stronger forecasting baseline.')
h('1.4 Why ML is appropriate')
p('Repeated sales histories provide supervised examples. A global tree-based regressor can learn nonlinear relationships between historical sales, planned prices, calendar effects and product/store attributes, while sharing information across series. This is useful when a single average cannot represent promotional uplift and weekday patterns. ML remains conditional on evidence: if it cannot beat a tuned seasonal baseline or produce business benefit in the pilot, the retailer retains the simpler policy.')
p('The forecast is an estimate of demand under adequate stock availability. Observed sales are censored during stockouts, so indiscriminately treating every recorded sale as true demand would reinforce under-ordering. Section 2 defines exclusions and evaluation safeguards for this limitation.')

page('1.5 Stakeholders and scope boundaries')
table(['Stakeholder','Goal / concern','Design response'],[
['Store planners and managers','Avoid empty shelves; understand proposals; retain local control.','Dashboard with reason codes, approved forecast version, explicit fallback status, override and approval.'],
['Commercial / finance owner','Improve sales availability without increasing inventory or operating costs.','Pilot business metrics, inventory guardrails and monthly operating cost ceiling.'],
['Data scientist / ML engineer','Reliable labels, fair comparisons and reproducible releases.','Time-based backtests, versioned datasets, cohort metrics and promotion gates.'],
['Data engineering / operations','Meet ordering cutoff; recover from failed jobs; maintain source contracts.','Daily quality gates, retry/checkpointing, atomic publication and documented recovery.'],
['IT security / governance','Limit data access, exposure and retention.','No customer identifiers; service roles, encryption and auditable access.'],
['Procurement / suppliers','Stable feasible orders within lead times and pack sizes.','ERP integration, capacity constraints and planner-reviewed exceptions.']],[115,173,195])
h('1.6 Conflicting goals and resolution')
p('Store managers favor extra safety stock, while finance wants less working capital. The solution sets an explicit stockout guardrail and evaluates inventory reduction only when availability is maintained. Planners may favor local adjustments, while the ML team needs consistent evaluation; both original forecasts and overridden orders are retained so forecast performance and decision performance can be reviewed independently.')
h('1.7 Scope and deployment context')
p('The MVP serves 40,000 active store-product pairs, generating 280,000 horizon rows per day. A single global model, one scheduled batch pipeline, a read-only forecast API and a planner dashboard are sufficient. The system runs in a managed cloud or existing company infrastructure with an encrypted object store and relational database. Training and inference run on CPU; specific resource sizes are established through load tests. Streaming ingestion, a separate feature-store service and Kubernetes are deferred until justified by measured constraints.')
h('1.8 Assumptions requiring validation')
p('The design assumes reliable daily POS exports, inventory snapshots, confirmed promotion plans, stable product IDs and support from an ERP integration owner. Discovery must validate store opening calendars, inventory accuracy, label availability, supplier lead time and comparable pilot stores. A sponsor must confirm whether the benefit warrants the operating cost. If stockout observations are unreliable, the pilot cannot claim a validated demand or stockout improvement until measurement is repaired.')

page('2. Requirements analysis (20%)')
h('2.1 Functional requirements and acceptance')
table(['ID','Requirement','Acceptance condition'],[
['FR-01','Predict nonnegative daily demand for horizons 1-7 for every eligible active store-product pair.','Seven unique dated outputs per eligible pair; exclusions receive seven explicit fallback rows.'],
['FR-02','Generate draft replenishment quantities using inventory position, lead time, safety stock and pack rules.','Order is nonnegative, respects pack size and has constraint/exception reasons.'],
['FR-03','Integrate daily POS, stock, product and promotion exports; export approved drafts to ERP.','Source totals reconcile; schema mismatch quarantined; idempotency key prevents duplicate ERP drafts.'],
['FR-04','Provide planner search/filter, forecast view, override and approval actions.','A planner can inspect a product, adjust a proposal with a reason and submit a versioned draft.'],
['FR-05','Store provenance and feedback.','Each row links source snapshot, feature/code/model versions, fallback flag and forecast origin; each approval records actor and time.'],
['FR-06','Recover safely from stale data, model failure or partial computation.','No partial run becomes active; affected pairs use marked fallback; failed integration is retried without duplication.']],[48,233,202])
h('2.2 Input specification')
p('Daily sales input: store_id, sku_id, local_business_date, units_sold, returns_units, unit_price and load_timestamp. Inventory: closing_on_hand, confirmed_inbound, reserved_units, stockout_flag and observed availability. Product master: category, pack_size, shelf_life_group and active dates. Store master: region and opening calendar. Confirmed promotion plan: effective dates, planned price and promotion type. ERP supplies supplier lead time and order caps. Units sold are gross nonnegative units; returns are separate and never turned into negative demand labels.')
h('2.3 Output specification and policy')
p('Each forecast row contains store_id, sku_id, forecast_origin, target_date, horizon_days, predicted_units, model_version, run_id, data_cutoff and quality_status. An order proposal adds inventory_position, lead_time_days, safety_stock_units, recommended_order_units, pack_size, exception_reason and approval_status. Example: origin 2026-10-03, target 2026-10-04, horizon 1, predicted_units 18.6; this is an illustrative schema value, not a model result.')
p('Inventory position = on-hand + confirmed inbound - reserved units. Draft order = round up to pack size(max(0, sum of forecasts over lead time plus one daily review interval + policy safety stock - inventory position)). Lead time plus review interval is at most four days in this scope. Safety stock initially uses an existing planner-approved rule. Shelf-life/storage and supplier caps may limit the order; a conflict raises a visible exception. Forecasts remain decimal estimates; ordering rounds units. ERP orders require planner approval.')

page('2.4 Non-functional and data requirements')
table(['Area','Proposed requirement','Verification / handling'],[
['Performance','Publish by 05:30 Asia/Bangkok daily; batch inference p95 <= 30 min at MVP load; cached API p95 <= 500 ms at 20 requests/s.','Measure job and API timings in representative load tests; do not equate API latency with model inference.'],
['Scale','40,000 active pairs; 280,000 daily rows. Design for 3x load (120,000 pairs; 840,000 rows) within 60 min inference.','Partition computation by store; use fixed feature snapshots and bounded parallel CPU workers.'],
['Reliability','>= 99% daily on-time complete publications over rolling 90 days; API >= 99.5% monthly availability.','Checkpoint jobs, atomic active-run pointer, two API replicas; fallback output allowed but quality status retained.'],
['Recovery','RTO <= 60 min from incident detection; RPO <= 24 h for durable source data.','Daily backup, replay immutable daily exports, restore drill quarterly; source cutoff misses trigger fallback.'],
['Maintainability','Weekly retraining candidate; daily data and job checks; weekly matured forecast review.','Version schemas/features/configuration; drift triggers investigation; all releases require quality gates and owner approval.'],
['Security','Role-based least privilege, encrypted transport/storage, audited approvals.','Separate ingestion/training/planner roles; secret manager; security review before deployment.']],[80,209,194])
h('2.5 Data sources, quality, volume and privacy')
p('Assume 18 months (approximately 548 days) of daily history: 20 x 2,000 x 548 = 21.92 million store-product-day rows. The raw daily exports arrive by 01:00. Allow approximately 100 GB for raw revisions, curated snapshots, artifacts and overhead; measure actual compressed size during discovery. Supervised multi-horizon expansion can reach approximately 153 million rows, so training samples origins deterministically or streams partitions rather than loading the full expansion into memory.')
p('Enforce uniqueness on store_id/sku_id/date, valid foreign keys, nonnegative gross sales and plausible inventory/price ranges. Detect missing files separately from genuine zero sales on an open day. Reconcile daily store sales totals to the POS source within 0.1%; expected eligible keys must be >= 99% complete and at least 99.5% of received rows must pass schema/range checks. Deduplicate by source revision and quarantine invalid rows. A failed store receives fallback outputs rather than fabricated zero-sales features. Missing critical inventory blocks its order export.')
p('Collect availability flags to identify censored labels. Exclude known stockout target-days from primary training and observed-sales error metrics; remove severely stock-constrained histories from ML eligibility. Lag features include availability/missingness flags so suppressed sales are visible. This reduces but does not remove selection bias: report availability coverage and validate final stockout outcomes prospectively. Never claim observed-sales error directly measures unobserved lost demand.')
p('Do not ingest customer names, loyalty IDs or basket-level identities. Daily aggregates suffice. Store employee identifiers appear only in restricted approval logs. Proposed retention: 24 months for raw/curated sales, 12 months for forecasts and approvals, and 90 days for operational logs, subject to organizational policy. Aggregate exports, restrict access, audit downloads and purge expired partitions and backups according to approved policy; no legal compliance claim is assumed.')

page('3. Goals and metrics (20%)')
h('3.1 Goal hierarchy and proposed targets')
p('Business goal: improve replenishment economics. System goal: deliver trustworthy proposals before ordering. Model goal: forecast available demand more accurately than seasonal rules. Model accuracy enables better decisions, but business improvement also depends on inventory accuracy, lead times and planner behavior. All targets below are proposed launch or pilot gates; baseline numbers are assumed until measured.')
table(['Layer / metric','Baseline','Target / decision rule'],[
['Business: stockout rate','Assumed 8% eligible open pair-days.','<= 6.4% during controlled pilot; 20% relative reduction.'],
['Business: inventory days of supply (DOS)','Assumed 12 days.','<= 10.8 days, with stockout guardrail maintained.'],
['Business: planner effort','Assumed 4 h/store/week.','<= 2.5 h/store/week, from time logs.'],
['Business: operating cost and benefit','No deployed ML cost; current waste/lost-margin baseline unmeasured.','Proposed recurring system cost <= USD 500/month; estimated monthly incremental margin + avoided waste - incremental operating cost > 0.'],
['System: daily deadline SLO','Manual delivery is not logged.','>= 99% complete by 05:30 over 90 days; ML-mode pair coverage >= 95%.'],
['System: API / batch','No comparable production measurement.','API p95 <= 500 ms at 20 requests/s; MVP inference p95 <= 30 min; monthly API uptime >= 99.5%.'],
['Model: unit WAPE, days 1-7 pooled','Illustrative tuned seasonal baseline 30%.','<= 25%, AND >= 10% relative improvement over the actually measured strongest baseline.'],
['Model: bias and cohort guardrails','To be measured for each baseline.','Absolute signed bias <= 5%; no material store/category cohort > 2 percentage points worse WAPE than baseline.']],[162,125,196])
h('3.2 Business metric definitions')
p('Stockout rate = eligible open store-product-days with any confirmed stockout / all eligible open store-product-days. Use consistent inventory observations in both arms. DOS = average inventory units / average daily units sold over the same evaluation window, calculated within SKU then aggregated with predeclared inventory-value weights; exclude zero-sales SKUs and report them separately. Stockouts can suppress the sales denominator, so interpret DOS alongside availability, inventory value and waste.')
p('Planner effort includes reviewing, overriding and submitting replenishment proposals, not merely time using the dashboard. Log representative weeks in both arms. Net benefit uses incremental gross margin and avoided spoilage measured against control, subtracting incremental hosting and support cost; avoid counting inventory cash release as recurring profit. Monetary targets remain provisional until finance supplies unit margins and holding costs.')

page('3.3 Model evaluation and business validation')
h('Metric formulas and interpretation')
p('For actual observed units y and prediction f on qualified target-days, WAPE = 100 x sum(|y - f|) / sum(y). This is unit-volume weighted across rows, not the average of individual percentage errors. If a cohort has zero total actual units, WAPE is undefined; report its MAE instead. Signed bias = 100 x sum(f - y) / sum(y); positive means overforecasting. Also report MAE in units, WAPE for each horizon and store/category slices. Define material cohorts before evaluation (>= 1% of evaluated sales volume and >= 100 qualified rows).')
p('Report stockout/missing-data exclusions and coverage next to every score. The same target rows, input cutoff and rounding conventions apply to candidate and baseline. Both baselines use data available at each historical origin: the current 28-day-average policy and the four-prior-same-weekday average. Tune simple promotion factors on training data only; compare to the best validation-selected baseline.')
h('Time-aware validation and leakage prevention')
p('Use an 18-month hypothetical timeline: months 1-12 train, month 13 validate; months 1-13 train, month 14 validate; months 1-14 train, month 15 validate. Freeze feature definitions, hyperparameters and baseline choice after these rolling folds. Refit through month 15, then evaluate once on untouched months 16-18. Group splits by forecast-origin date across all stores. Ensure each training origin plus its seven-day target window ends before the first validation/test origin, and restrict all feature records by their historical availability timestamp. This prevents overlapping future labels and revision leakage. Time-based validation follows the ordering principle described in scikit-learn documentation [2].')
p('Use direct multi-horizon supervised learning: one global model consumes origin features plus horizon and known target-day calendar/promotion attributes. Sales lags 1/7/14/28 and shifted rolling windows use only dates before the origin cutoff. Do not use actual future sales, realized future prices or retrospectively revised promotions. At an origin, future planned prices are used only if confirmed before the cutoff. Stockout target masks affect eligibility, never become future-known features.')
h('Pilot design and release decisions')
p('Run two weeks in shadow mode, then an eight-week pilot with 10 matched store pairs randomized to treatment/control, blocking on size and category mix. Treatment planners receive ML drafts; control keeps the current policy. Record overrides, promotion exposure, supplier failures and demand censoring. Use store-clustered uncertainty intervals and predeclare the primary stockout metric, DOS guardrail and exclusion rules. Historical variability and the business minimum effect determine power before launch; extend the pilot if 20 stores and eight weeks are insufficient.')
p('Business launch requires the stockout point target and an uncertainty interval supporting improvement, a maintained inventory/waste guardrail, acceptable planner effort and positive estimated net benefit. If intervals remain inconclusive, continue measurement rather than claim success. Passing offline gates permits shadow/pilot deployment only. Finance and operations approve broader rollout based on pilot evidence; forecast errors alone cannot authorize it.')

page('4. High-level architecture design (25%)')
h('4.1 Architecture and data/control flow')
story.append(Architecture())
diagram_paragraph=doc.add_paragraph()
md.append('![Production architecture](architecture.png)\n')
h('4.2 Architectural rationale')
p('Daily order decisions do not need streaming predictions. Durable raw exports and one shared feature-building library keep training and batch inference consistent. The registry stores evaluated model artifacts; inference pins an approved version for each run rather than accepting a mutable alias midway through a job. A forecast database isolates planner API traffic from expensive feature construction and training. Atomic publication prevents planners from seeing a partially populated day.')
p('The diagram separates training from operational inference. Monitoring receives delayed actuals and order outcomes; a drift alert creates a retraining request, not an automatic production swap. The replenishment step is deterministic business logic using one demand model, which keeps complexity within the assignment scope.')
h('4.3 Daily timeline (Asia/Bangkok)')
table(['Time','Activity / responsibility'],[['01:00','Daily source cutoff; data engineer owns export completeness.'],['01:00-02:00','Ingest, reconcile, validate; one automatic source retry and incident alert.'],['02:00-03:00','Build origin-safe features; freeze dataset snapshot.'],['03:00-03:30','Score seven horizons; validate outputs and compute draft orders.'],['03:30-05:30','Publish or recover; use explicit fallback for unavailable pairs; notify planners.'],['06:00 onward','Planner review and approval; ERP draft submission using idempotency keys.']],[95,388])

page('4.4 Components and ML pipeline stages')
table(['Component / tools considered','Purpose and responsibility'],[
['Ingestion: scheduled Python jobs, ERP/POS exports','Fetch revision-aware files; validate schemas/keys and reconcile totals; retain immutable source manifests.'],
['Storage: encrypted object store + Parquet','Separate raw and curated areas; partition by business date/store; snapshots and checksums support replay.'],
['Feature processing: Python + DuckDB / pandas','Build lags, rolling statistics, calendar/planned promotion features and availability flags; identical code for historical and daily origins.'],
['Training: LightGBM on CPU','One global gradient-boosted regressor with a nonnegative count objective (Poisson) [1]; validate hyperparameters, clip tiny numerical negatives and compare baselines.'],
['Experiments / registry: MLflow','Record code, features, dataset manifest, parameters and scores; register candidate; approval-controlled model alias and immutable version pin [3].'],
['Orchestration: managed scheduler or existing Airflow','Daily tasks and weekly candidate training; retries, dependencies, checkpoints and run state. Prefer an existing platform to reduce support work.'],
['Serving: PostgreSQL + FastAPI','Store run-versioned forecasts and proposals; indexed read API, authorization, active-run pointer and approval records.'],
['UI / integration: internal dashboard + ERP API','Explain inputs and reasons, permit auditable overrides; submit approved drafts once, reconcile ERP acknowledgment.'],
['Monitoring: Prometheus/Grafana or existing equivalent','Job/deadline/error alerts, source health, weekly matured quality and business outcome dashboards.']],[179,304])
h('4.5 Training and release pipeline')
p('Ingestion creates immutable raw records and a source manifest. Validation produces a curated daily table and quarantined exceptions. Feature generation expands forecast origins by horizon and joins origin-valid data. Training excludes censored target-days, records sample eligibility and deterministically samples origin dates if necessary. Validation reports baseline comparisons, bias, cohort quality and compute cost. MLflow tracks artifacts and model versions; a reviewer checks reproducibility and all gates before assigning the production alias. Inference resolves that alias once and records its immutable version in the run [3].')
p('CI checks schema contracts, temporal joins, stockout masks, identical historical/daily feature logic and order rounding. A release bundles model, feature schema and configuration. Reject promotion when columns, categorical mappings or source semantics are incompatible. Re-running a historical origin must use archived inputs and the appropriate artifact, not today\'s revised master data. Containerized dependencies and code hashes allow reproducible training and rollback.')

page('4.6 Reliability, monitoring and maintenance')
h('Inference, publication and integration')
p('Daily inference reads one frozen input snapshot and scores by store partition. Persist each partition under a run_id; resume failed partitions without overwriting published results. Validate expected key/horizon counts, uniqueness, finite/nonnegative predictions, model version and outlier checks. Build order proposals against the matching inventory snapshot. Commit a complete run and switch the active-run pointer in a single database transaction. Planner approvals retain the viewed run_id so a later refresh cannot silently change an approved order.')
p('ERP submissions use an idempotency key combining store, supplier, review date and approved revision. If an acknowledgment times out, query ERP by key before retrying. Corrections create explicit superseding drafts rather than duplicate orders. Product or price drift does not bypass supplier constraints.')
h('Failure modes and graceful degradation')
table(['Failure','Operational response'],[
['POS missing / data quality gate fails','Retry within deadline, then flag affected store/pairs. Recompute seasonal fallback with last valid history and current confirmed calendar. Display age and fallback status.'],
['Model scoring or registry unavailable','Use cached approved artifact if compatible; otherwise deterministic seasonal fallback. If historical inputs are too old (> 48 h), show unavailable and route to manual planning.'],
['Inventory snapshot missing / stale','Forecast may remain viewable; block order export until inventory is repaired or planner follows documented manual ERP process.'],
['API/database outage','Fail over API replica and restore database; export signed/versioned contingency CSV for planners. Record incident and recover within RTO.'],
['Candidate regression / production anomaly','Retain current model; rollback model+feature/config bundle, re-score unpublished drafts and notify planners. Approved orders require explicit reconciliation.']],[150,333])
h('Monitoring triggers and ownership')
p('Operations receives alerts for missed 03:30 expected publication, forecast/API error spikes, source-quality failures and deadline risk; acknowledge critical alerts within 15 minutes. Data engineering investigates completeness below 99% or received-row validity below 99.5%. ML engineering reviews weekly matured WAPE above 25%, absolute bias above 5%, a > 10% relative WAPE deterioration against baseline for two windows, or material feature distribution changes. Reference drift uses training distributions and fixed bins; a population stability index above 0.2 is a heuristic investigation signal, not proof that accuracy has failed.')
p('Calculate horizon-seven scores only after all target actuals arrive. Aggregate data quality and fallback coverage daily; compare override rates, order outcomes and stockouts weekly. Record model eligibility and excluded cohorts, so attractive aggregate scores do not hide poor coverage. Weekly retraining creates a candidate; promotion requires the same gates as initial release. Quarterly restore drills and dependency reviews keep recovery and maintenance credible.')

page('5. Trade-offs analysis (15%)')
table(['Trade-off','Options and consequence','Decision and revisit condition'],[
['Accuracy vs. latency','Seasonal rules are cheap and transparent but miss nonlinear promotion effects. A global boosted tree can improve error with additional feature and training cost; a deep sequence model adds compute and tuning burden.','Use one LightGBM model on CPU and retain seasonal fallback. Reject extra complexity unless measured WAPE/cohort gains and pilot benefits justify it while meeting batch timing.'],
['Freshness vs. cost','Streaming inventory and sales could respond intraday but requires event infrastructure and operational support. Daily snapshots may miss changes after the cutoff.','Use daily batch forecasts because orders are reviewed daily. Display cutoff, honor confirmed inventory at approval, and investigate intraday refresh if late demand shocks consistently harm outcomes.'],
['Simplicity vs. performance','One global model shares data and is easy to release. Separate per-product/category models may fit local patterns but multiply artifacts, sparse-series issues and maintenance.','Use a global model with category/store features and cohort gates; exclude unsupported cohorts. Revisit segmentation only if persistent material cohort regressions cannot be fixed within the single model.'],
['Automation vs. control','Automatic ERP orders reduce labor but can amplify bad data or demand spikes. Human approval adds work and variable intervention.','Produce draft orders only; require reasons for overrides and expose fallback/constraint alerts. Consider bounded automatic approval after sustained pilot performance and explicit business authorization.'],
['Privacy vs. granularity','Customer/basket histories may add behavioral signal but increase identity exposure and governance burden. Aggregate daily sales loses individual patterns.','Use store-product aggregates only; no personalization is needed for replenishment. Consider additional aggregate signals after a documented accuracy/cost review, without assuming identity data is required.']],[97,193,193])
h('5.1 Residual risks and decision discipline')
p('The main uncertainty is censored demand: stockout exclusion can bias training toward well-stocked days. Inventory reconciliation and the controlled business pilot address this more credibly than invented lost-sales labels. Cold starts and sparse products remain fallback-only. Abrupt promotional or supplier changes can exceed the historical training distribution, so constraints, freshness indicators and human review remain essential. Large overrides may improve orders while reducing interpretability of model-to-business attribution; retain both forecasts and decisions.')
p('A USD 500 monthly budget is a proposed constraint, not a cloud pricing estimate. Benchmark storage, CPU hours, database availability and monitoring against local procurement quotes, and separately account for support labor. Adopt only features that meet their acceptance criteria and demonstrate value. If the baseline wins or the pilot remains inconclusive, preserve the current policy and document the evidence rather than force an ML rollout.')

page('References and requirement traceability')
h('References')
p('[A] Assignment brief: INDIVIDUAL ASSIGNMENT 1_baiCaNhan.pdf, supplied locally, 3 pages. Sections I-IV define the scenario characteristics, report rubric and submission naming.')
p('[1] LightGBM, official Parameters documentation, objective options including Poisson. https://lightgbm.readthedocs.io/en/latest/Parameters.html (accessed 3 October 2026). Supports the candidate algorithm capability; does not substantiate performance claims.')
p('[2] scikit-learn, official TimeSeriesSplit documentation. https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.TimeSeriesSplit.html (accessed 3 October 2026). Supports chronological validation and the gap concept; grouped origin/label embargo rules here are scenario-specific design decisions.')
p('[3] MLflow, official Model Registry Workflows documentation. https://www.mlflow.org/docs/latest/ml/model-registry/workflow/ (accessed 3 October 2026). Supports model version/alias lifecycle and approval-driven artifact selection; operational gates here are proposed requirements.')
h('Completeness against the PDF')
table(['PDF requirement','Report coverage'],[
['Scenario: ML task; >= 2 stakeholders; real deployment; realistic data; measurable impact','Sections 1.1-1.8, 2.5 and 3.1-3.3; hypothetical retail scenario with six stakeholder groups and one ML model.'],
['Problem definition (20%): context, statement, current handling, justification, stakeholders','Sections 1.1-1.8.'],
['Requirements analysis (20%): functional, I/O, integration, user interaction','Sections 2.1-2.3; planner approval, ERP contracts and provenance.'],
['Requirements: latency, scale, reliability, maintainability; sources, quality, privacy, volume','Sections 2.4-2.5 and 4.6.'],
['Goals/metrics (20%): business, system, model, thresholds and baselines','Sections 3.1-3.3, with formulas, assumed baselines, release gates and pilot design.'],
['Architecture (25%): clear diagram, data flow, pipeline stages, component roles and tools','Sections 4.1-4.6: ingestion, preprocessing, training, inference, monitoring, recovery.'],
['Trade-offs (15%): at least 3-4 significant decisions','Section 5 analyzes five trade-offs, decisions and conditions for revisiting them.'],
['Submission naming: DDM501_Assignment1_[StudentID]_[Name].pdf','Generated filename follows the required template; replace placeholders with verified student information.']],[213,270])
h('Submission note')
p('This is a design document, not an implemented or experimentally validated system. All assumed values must remain labeled as assumptions if used for submission. Student identifiers must be completed before final submission. The editable Word and Markdown copies are provided for review; the PDF is the required submission format.')

def footer(c,d):
    c.saveState();c.setStrokeColor(colors.HexColor('#cbd9e1'));c.line(56,43,A4[0]-56,43)
    c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#536b80'))
    c.drawString(56,30,'DDM501 | Demand Forecasting | ML System Design')
    c.drawRightString(A4[0]-56,30,str(d.page));c.restoreState()

pdf=OUT/(stem+'.pdf')
SimpleDocTemplate(str(pdf),pagesize=A4,rightMargin=56,leftMargin=56,topMargin=48,bottomMargin=58,title='DDM501 Assignment 1 - Retail Demand Forecasting',author=profile['name']).build(story,onFirstPage=footer,onLaterPages=footer)
(OUT/'REPORT.md').write_text('\n'.join(md),encoding='utf-8')
# Standalone vector architecture plus a raster preview for the editable document.
diagram_pdf=OUT/'architecture.pdf'
c=canvas.Canvas(str(diagram_pdf),pagesize=(503,340));c.translate(10,10);Architecture().drawOn(c,0,0);c.save()
with pymupdf.open(diagram_pdf) as rendered:
    rendered[0].get_pixmap(matrix=pymupdf.Matrix(2,2)).save(str(OUT/'architecture.png'))
diagram_paragraph.add_run().add_picture(str(OUT/'architecture.png'),width=Inches(6.2))
doc.save(OUT/(stem+'.docx'))
(OUT/'architecture.mmd').write_text('''flowchart LR
  Sources[POS / Inventory / ERP / Promotion plans] --> Ingest[Daily ingestion and validation]
  Ingest --> Raw[Immutable raw snapshots]
  Raw --> Features[Versioned curated features]
  Features --> Train[Weekly training and time backtests]
  Train --> Gate[Evaluation and human promotion gate]
  Gate --> Registry[MLflow registry: pinned version]
  Features --> Batch[Daily batch inference]
  Registry --> Batch
  Batch --> Policy[Forecast validation and replenishment rules]
  Policy --> DB[Atomic publication / PostgreSQL]
  DB --> API[Read API and planner dashboard]
  API --> Approval[Planner approval / override]
  Approval --> ERP[Idempotent ERP draft orders]
  Sources --> Monitor[Quality / outcomes / drift monitoring]
  Batch --> Monitor
  Approval --> Monitor
  Monitor --> Review[Investigation / retraining request]
  Review --> Train
''',encoding='utf-8')
print(f'Created {pdf.name}, editable DOCX, REPORT.md and architecture assets.')
